import tkinter as tk
from tkinter import filedialog, messagebox
import threading
import time
import os

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright


# =========================================================
# SETTINGS
# =========================================================

EDUCUBE_URL = "https://cms.educube.net/cms/Curriculum/ClassPlanning"

VALID_HOUSES = {
    "hope",
    "peace",
    "unity",
    "love"
}

DEFAULT_DELAY = 1.0

SHORT_TIMEOUT = 1500
NORMAL_TIMEOUT = 5000


# =========================================================
# EVENTS
# =========================================================

start_event = threading.Event()
stop_event = threading.Event()
close_browser_event = threading.Event()


# =========================================================
# TEXT HELPERS
# =========================================================

def clean_text(value):
    if value is None:
        return ""

    return " ".join(str(value).strip().split())


def normalize(value):
    return clean_text(value).lower()


def normalize_gr(value):
    """
    Keeps GR Number exactly as text.

    Examples:
        76/96P
        03/98P
        35/140PP
    """

    if value is None:
        return ""

    text = str(value).strip()

    # Remove accidental Excel numeric .0
    if text.endswith(".0"):
        text = text[:-2]

    # Remove unnecessary spaces
    text = " ".join(text.split())

    return text


# =========================================================
# READ EXCEL
# =========================================================

def read_excel(file_path):

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    required_headers = {
        "gr number",
        "house name"
    }

    selected_sheet = None
    header_row = None
    column_map = None

    # -----------------------------------------------------
    # Automatically find correct sheet and header row
    # -----------------------------------------------------

    for sheet in workbook.worksheets:

        max_check_row = min(
            sheet.max_row,
            20
        )

        for row_number in range(
            1,
            max_check_row + 1
        ):

            headers = {}

            for col_number in range(
                1,
                sheet.max_column + 1
            ):

                value = sheet.cell(
                    row=row_number,
                    column=col_number
                ).value

                if value is not None:

                    headers[
                        normalize(value)
                    ] = col_number

            if required_headers.issubset(
                set(headers.keys())
            ):

                selected_sheet = sheet
                header_row = row_number

                column_map = {
                    "gr": headers["gr number"],
                    "house": headers["house name"]
                }

                break

        if selected_sheet is not None:
            break

    # -----------------------------------------------------
    # Excel structure not found
    # -----------------------------------------------------

    if selected_sheet is None:

        raise ValueError(
            "Could not find the correct Excel sheet.\n\n"
            "Your Excel must contain these headers:\n\n"
            "GR Number\n"
            "House Name"
        )

    # -----------------------------------------------------
    # Print Excel information
    # -----------------------------------------------------

    print()
    print("==========================================")
    print("EXCEL INFORMATION")
    print("==========================================")

    print(
        "Using sheet       :",
        selected_sheet.title
    )

    print(
        "Header row        :",
        header_row
    )

    print(
        "GR Number col     :",
        column_map["gr"]
    )

    print(
        "House Name col    :",
        column_map["house"]
    )

    print("==========================================")
    print()

    # -----------------------------------------------------
    # Read students
    # -----------------------------------------------------

    students = []

    seen_gr_numbers = set()

    for row_number in range(
        header_row + 1,
        selected_sheet.max_row + 1
    ):

        gr_number = selected_sheet.cell(
            row=row_number,
            column=column_map["gr"]
        ).value

        house = selected_sheet.cell(
            row=row_number,
            column=column_map["house"]
        ).value

        gr_number = normalize_gr(gr_number)
        house = clean_text(house)

        # Blank GR
        if not gr_number:
            continue

        # Blank House
        if not house:

            print(
                f"Row {row_number}: "
                f"Blank house for GR {gr_number}"
            )

            continue

        # Invalid House
        if normalize(house) not in VALID_HOUSES:

            print(
                f"Row {row_number}: "
                f"Invalid house '{house}' "
                f"for GR {gr_number}"
            )

            continue

        # Duplicate GR in Excel
        if gr_number in seen_gr_numbers:

            print(
                f"Row {row_number}: "
                f"Duplicate GR Number "
                f"{gr_number} skipped"
            )

            continue

        seen_gr_numbers.add(gr_number)

        students.append(
            {
                "excel_row": row_number,
                "gr": gr_number,
                "house": house
            }
        )

    return students, selected_sheet.title


# =========================================================
# FIND GR NUMBER INPUT
# =========================================================

def find_gr_input(page, gr_number):

    gr_number = normalize_gr(gr_number)

    print()
    print(
        f"Searching GR Number: {gr_number}"
    )

    try:

        # -------------------------------------------------
        # IMPORTANT:
        # GR NUMBER IS AN INPUT VALUE, NOT TEXT.
        # -------------------------------------------------

        inputs = page.locator(
            'input[name="gr_number"]'
        )

        count = inputs.count()

        print(
            f"GR input fields found: {count}"
        )

        for i in range(count):

            try:

                input_box = inputs.nth(i)

                value = input_box.input_value(
                    timeout=SHORT_TIMEOUT
                )

                value = normalize_gr(value)

                if value == gr_number:

                    print(
                        f"✓ GR Number matched: "
                        f"{gr_number}"
                    )

                    return input_box

            except Exception:
                continue

    except Exception as e:

        print(
            "Error searching GR input:",
            e
        )

    print(
        f"GR {gr_number}: NOT FOUND"
    )

    return None


# =========================================================
# FIND STUDENT ROW
# =========================================================

def find_student_row(page, gr_number):

    gr_input = find_gr_input(
        page,
        gr_number
    )

    if gr_input is None:
        return None

    try:

        row = gr_input.locator(
            "xpath=ancestor::tr[1]"
        )

        if row.count() == 0:

            print(
                "GR input found but row not found."
            )

            return None

        # -------------------------------------------------
        # Print row information
        # -------------------------------------------------

        try:

            row_text = row.inner_text(
                timeout=SHORT_TIMEOUT
            )

            print(
                "Matched row:",
                clean_text(row_text)
            )

        except Exception:
            pass

        return row

    except Exception as e:

        print(
            "Error finding student row:",
            e
        )

        return None


# =========================================================
# FIND HOUSE DROPDOWN
# =========================================================

def find_house_dropdown(row):

    try:

        # -------------------------------------------------
        # EduCube specifically uses:
        #
        # select[name="house_name"]
        # -------------------------------------------------

        dropdown = row.locator(
            'select[name="house_name"]'
        )

        if dropdown.count() > 0:

            return dropdown.first

    except Exception as e:

        print(
            "House dropdown search error:",
            e
        )

    # -----------------------------------------------------
    # Fallback: inspect every select
    # -----------------------------------------------------

    try:

        selects = row.locator("select")

        count = selects.count()

        print(
            f"Fallback dropdown count: {count}"
        )

        for i in range(count):

            try:

                select = selects.nth(i)

                name = select.get_attribute(
                    "name"
                )

                if name == "house_name":

                    return select

            except Exception:
                continue

    except Exception:
        pass

    return None


# =========================================================
# CHANGE HOUSE
# =========================================================

def change_house(row, house):

    house = clean_text(house)

    print()
    print(
        f"Changing House to: {house}"
    )

    dropdown = find_house_dropdown(
        row
    )

    if dropdown is None:

        return (
            False,
            "House dropdown not found"
        )

    # -----------------------------------------------------
    # Wait until dropdown is available
    # -----------------------------------------------------

    try:

        dropdown.wait_for(
            state="visible",
            timeout=NORMAL_TIMEOUT
        )

    except Exception as e:

        return (
            False,
            f"House dropdown not visible: {e}"
        )

    # -----------------------------------------------------
    # Read all options
    # -----------------------------------------------------

    try:

        options = dropdown.locator(
            "option"
        )

        option_count = options.count()

        print(
            f"House options found: "
            f"{option_count}"
        )

        target_option_value = None

        for i in range(option_count):

            option = options.nth(i)

            try:

                option_text = clean_text(
                    option.inner_text(
                        timeout=SHORT_TIMEOUT
                    )
                )

                option_value = option.get_attribute(
                    "value"
                )

                print(
                    f"  Option {i + 1}: "
                    f"text='{option_text}' "
                    f"value='{option_value}'"
                )

                if normalize(
                    option_text
                ) == normalize(house):

                    target_option_value = option_value

            except Exception:
                continue

        # -------------------------------------------------
        # House option does not exist
        # -------------------------------------------------

        if target_option_value is None:

            return (
                False,
                f"House option '{house}' "
                f"not found in dropdown"
            )

        # -------------------------------------------------
        # METHOD 1:
        # Select using exact OPTION VALUE
        # -------------------------------------------------

        try:

            dropdown.select_option(
                value=target_option_value,
                timeout=NORMAL_TIMEOUT
            )

            print(
                f"✓ House changed to {house}"
            )

            return (
                True,
                f"House changed to {house}"
            )

        except Exception as e:

            print(
                "Normal select_option failed:",
                e
            )

        # -------------------------------------------------
        # METHOD 2:
        # JavaScript fallback
        #
        # This is useful if EduCube has special
        # behaviour around the select element.
        # -------------------------------------------------

        try:

            dropdown.evaluate(
                """
                (select, targetValue) => {

                    select.value = targetValue;

                    select.dispatchEvent(
                        new Event(
                            'change',
                            {
                                bubbles: true
                            }
                        )
                    );

                    select.dispatchEvent(
                        new Event(
                            'input',
                            {
                                bubbles: true
                            }
                        )
                    );
                }
                """,
                target_option_value
            )

            # Verify
            current_value = dropdown.input_value()

            if current_value == target_option_value:

                print(
                    f"✓ House changed to {house} "
                    f"(JavaScript method)"
                )

                return (
                    True,
                    f"House changed to {house}"
                )

        except Exception as e:

            print(
                "JavaScript selection failed:",
                e
            )

    except Exception as e:

        print(
            "House option inspection failed:",
            e
        )

    return (
        False,
        f"Could not change House to {house}"
    )


# =========================================================
# AUTOMATION
# =========================================================

def run_automation(
    file_path,
    delay,
    log_function
):

    try:

        students, sheet_name = read_excel(
            file_path
        )

    except Exception as e:

        log_function("")
        log_function(
            "ERROR READING EXCEL:"
        )
        log_function(str(e))

        return

    if not students:

        log_function("")
        log_function(
            "No valid GR Number records found."
        )

        return

    # -----------------------------------------------------
    # GUI LOG
    # -----------------------------------------------------

    log_function("")
    log_function(
        f"Excel sheet detected: {sheet_name}"
    )

    log_function(
        f"Valid GR records loaded: "
        f"{len(students)}"
    )

    # -----------------------------------------------------
    # START PLAYWRIGHT
    # -----------------------------------------------------

    with sync_playwright() as p:

        browser = p.chromium.launch(
            headless=False
        )

        context = browser.new_context(
            viewport={
                "width": 1400,
                "height": 900
            }
        )

        page = context.new_page()

        log_function("")
        log_function(
            "Opening EduCube..."
        )

        try:

            page.goto(
                EDUCUBE_URL,
                wait_until="domcontentloaded",
                timeout=30000
            )

        except Exception as e:

            log_function(
                "Page loading warning: "
                + str(e)
            )

        log_function(
            "EduCube opened."
        )

        log_function("")
        log_function(
            "LOGIN MANUALLY."
        )

        log_function(
            "Open Curriculum → Class Planning."
        )

        log_function(
            "Make sure the complete student list "
            "is visible."
        )

        log_function("")
        log_function(
            "Then click START AUTOMATION."
        )

        # -------------------------------------------------
        # WAIT FOR START
        # -------------------------------------------------

        while not start_event.is_set():

            if stop_event.is_set():

                browser.close()
                return

            time.sleep(0.2)

        start_event.clear()

        # -------------------------------------------------
        # START MESSAGE
        # -------------------------------------------------

        log_function("")
        log_function(
            "=========================================="
        )

        log_function(
            "GR NUMBER AUTOMATION STARTED"
        )

        log_function(
            "Matching ONLY by GR Number"
        )

        log_function(
            "Class is NOT used"
        )

        log_function(
            "Section is NOT used"
        )

        log_function(
            "Student Name is NOT used"
        )

        log_function(
            "Save / Submit will NOT be clicked"
        )

        log_function(
            f"Delay: {delay} second(s)"
        )

        log_function(
            "=========================================="
        )

        # -------------------------------------------------
        # COUNTERS
        # -------------------------------------------------

        successful = 0
        not_found = 0
        errors = 0

        total = len(students)

        # -------------------------------------------------
        # PROCESS STUDENTS
        # -------------------------------------------------

        for index, student in enumerate(
            students,
            start=1
        ):

            if stop_event.is_set():

                log_function("")
                log_function(
                    "Automation stopped by user."
                )

                break

            gr = student["gr"]
            house = student["house"]

            log_function(
                f"[{index}/{total}] "
                f"GR: {gr} → {house}"
            )

            try:

                # -----------------------------------------
                # Find row using GR input value
                # -----------------------------------------

                row = find_student_row(
                    page,
                    gr
                )

                if row is None:

                    log_function(
                        "    ✗ GR Number not found"
                    )

                    not_found += 1

                    continue

                log_function(
                    "    ✓ Correct GR row found"
                )

                # -----------------------------------------
                # Change House
                # -----------------------------------------

                success, message = change_house(
                    row,
                    house
                )

                if success:

                    log_function(
                        "    ✓ " + message
                    )

                    successful += 1

                else:

                    log_function(
                        "    ✗ " + message
                    )

                    errors += 1

            except Exception as e:

                log_function(
                    "    ✗ ERROR: "
                    + str(e)
                )

                errors += 1

            # ---------------------------------------------
            # Delay
            # ---------------------------------------------

            time.sleep(
                max(1.0, delay)
            )

        # -------------------------------------------------
        # FINISHED
        # -------------------------------------------------

        log_function("")
        log_function(
            "=========================================="
        )

        log_function(
            "AUTOMATION FINISHED"
        )

        log_function(
            "=========================================="
        )

        log_function(
            f"Successful : {successful}"
        )

        log_function(
            f"Not Found  : {not_found}"
        )

        log_function(
            f"Errors     : {errors}"
        )

        log_function(
            "=========================================="
        )

        log_function("")
        log_function(
            "Browser is still OPEN."
        )

        log_function(
            "Please manually RECHECK all changes."
        )

        log_function(
            "SAVE / SUBMIT was NOT clicked."
        )

        # -------------------------------------------------
        # Keep browser open
        # -------------------------------------------------

        while not close_browser_event.is_set():

            time.sleep(0.5)

        browser.close()


# =========================================================
# TKINTER GUI
# =========================================================

root = tk.Tk()

root.title(
    "CMS EduCube GR Number House Auto Filler"
)

root.geometry(
    "900x680"
)

root.minsize(
    750,
    550
)


# =========================================================
# VARIABLES
# =========================================================

excel_path = tk.StringVar()

delay_value = tk.DoubleVar(
    value=1.0
)


# =========================================================
# LOG FUNCTION
# =========================================================

def log(message):

    def update():

        log_box.config(
            state=tk.NORMAL
        )

        log_box.insert(
            tk.END,
            message + "\n"
        )

        log_box.see(
            tk.END
        )

        log_box.config(
            state=tk.DISABLED
        )

    root.after(
        0,
        update
    )


# =========================================================
# BROWSE EXCEL
# =========================================================

def browse_excel():

    file_path = filedialog.askopenfilename(
        title="Select GR Number House Excel",
        filetypes=[
            (
                "Excel Files",
                "*.xlsx *.xlsm"
            )
        ]
    )

    if file_path:

        excel_path.set(
            file_path
        )

        log(
            "Selected Excel: "
            + os.path.basename(file_path)
        )


# =========================================================
# OPEN EDUCUBE
# =========================================================

def open_educube():

    file_path = excel_path.get()

    if not file_path:

        messagebox.showwarning(
            "Excel Required",
            "Please select your Excel file first."
        )

        return

    if not os.path.exists(
        file_path
    ):

        messagebox.showerror(
            "File Not Found",
            "Selected Excel file does not exist."
        )

        return

    stop_event.clear()
    start_event.clear()
    close_browser_event.clear()

    # Clear log
    log_box.config(
        state=tk.NORMAL
    )

    log_box.delete(
        "1.0",
        tk.END
    )

    log_box.config(
        state=tk.DISABLED
    )

    try:

        delay = float(
            delay_value.get()
        )

    except Exception:

        delay = 1.0

    delay = max(
        1.0,
        delay
    )

    threading.Thread(
        target=run_automation,
        args=(
            file_path,
            delay,
            log
        ),
        daemon=True
    ).start()


# =========================================================
# START
# =========================================================

def start_automation():

    start_event.set()


# =========================================================
# STOP
# =========================================================

def stop_automation():

    stop_event.set()

    log(
        "Stopping automation..."
    )


# =========================================================
# CLOSE
# =========================================================

def close_program():

    close_browser_event.set()

    root.destroy()


# =========================================================
# TITLE
# =========================================================

tk.Label(
    root,
    text=(
        "CMS EduCube GR Number House Auto Filler"
    ),
    font=(
        "Arial",
        20,
        "bold"
    )
).pack(
    pady=(15, 5)
)


tk.Label(
    root,
    text=(
        "GR Number → House\n"
        "GR Number is the ONLY student identifier\n"
        "Save / Submit is NEVER clicked automatically"
    ),
    font=(
        "Arial",
        10
    )
).pack(
    pady=(0, 15)
)


# =========================================================
# EXCEL FRAME
# =========================================================

excel_frame = tk.Frame(
    root
)

excel_frame.pack(
    fill="x",
    padx=20,
    pady=5
)


tk.Label(
    excel_frame,
    text="Excel File:",
    font=(
        "Arial",
        10,
        "bold"
    )
).pack(
    side="left"
)


tk.Entry(
    excel_frame,
    textvariable=excel_path,
    width=70
).pack(
    side="left",
    padx=10,
    fill="x",
    expand=True
)


tk.Button(
    excel_frame,
    text="Browse",
    command=browse_excel,
    width=10
).pack(
    side="left"
)


# =========================================================
# DELAY FRAME
# =========================================================

delay_frame = tk.Frame(
    root
)

delay_frame.pack(
    fill="x",
    padx=20,
    pady=8
)


tk.Label(
    delay_frame,
    text="Delay per student (seconds):"
).pack(
    side="left"
)


tk.Spinbox(
    delay_frame,
    from_=1.0,
    to=10.0,
    increment=1.0,
    textvariable=delay_value,
    width=8
).pack(
    side="left",
    padx=10
)


# =========================================================
# BUTTON FRAME
# =========================================================

button_frame = tk.Frame(
    root
)

button_frame.pack(
    pady=12
)


tk.Button(
    button_frame,
    text="OPEN EDUcube",
    command=open_educube,
    width=18,
    height=2
).grid(
    row=0,
    column=0,
    padx=5
)


tk.Button(
    button_frame,
    text="START AUTOMATION",
    command=start_automation,
    width=20,
    height=2
).grid(
    row=0,
    column=1,
    padx=5
)


tk.Button(
    button_frame,
    text="STOP",
    command=stop_automation,
    width=12,
    height=2
).grid(
    row=0,
    column=2,
    padx=5
)


# =========================================================
# INSTRUCTIONS
# =========================================================

instructions = (
    "1. Select the Excel file.\n"
    "2. Click OPEN EDUcube.\n"
    "3. Login manually.\n"
    "4. Open Curriculum → Class Planning.\n"
    "5. Make sure the student list is visible.\n"
    "6. Click START AUTOMATION.\n"
    "7. Automation finds the GR Number input field.\n"
    "8. It opens the House dropdown in that same row.\n"
    "9. It changes only that student's House.\n"
    "10. It does NOT click Save / Submit.\n"
    "11. Manually recheck and save."
)


tk.Label(
    root,
    text=instructions,
    justify="left",
    anchor="w",
    font=(
        "Arial",
        10
    )
).pack(
    fill="x",
    padx=30,
    pady=5
)


# =========================================================
# LOG FRAME
# =========================================================

log_frame = tk.Frame(
    root
)

log_frame.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=10
)


scrollbar = tk.Scrollbar(
    log_frame
)

scrollbar.pack(
    side="right",
    fill="y"
)


log_box = tk.Text(
    log_frame,
    height=18,
    width=100,
    font=(
        "Consolas",
        9
    ),
    yscrollcommand=scrollbar.set,
    state=tk.DISABLED
)

log_box.pack(
    side="left",
    fill="both",
    expand=True
)


scrollbar.config(
    command=log_box.yview
)


# =========================================================
# CLOSE EVENT
# =========================================================

root.protocol(
    "WM_DELETE_WINDOW",
    close_program
)


# =========================================================
# START GUI
# =========================================================

root.mainloop()